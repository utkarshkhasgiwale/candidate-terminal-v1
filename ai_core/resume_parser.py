import json
import os
import re
from dotenv import load_dotenv
from groq import Groq
from docx import Document
from pydantic import BaseModel
from pathlib import Path


# ============================================================
# Setup
# ============================================================

load_dotenv()

api_key = os.getenv("GROQ_API_KEY")

if not api_key:
    raise ValueError("API key not found")

client = Groq(api_key=api_key)

model = "openai/gpt-oss-120b"


# ============================================================
# Resume reading
# ============================================================

def read_docx(file_path: str) -> str:
    document = Document(file_path)
    text = ""

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"

    return text


# ============================================================
# Pydantic schemas
# ============================================================

class ExperienceEntry(BaseModel):
    role: str
    company: str
    duration: str
    description: str


class ProjectEntry(BaseModel):
    id: str
    name: str
    description: str
    tech: list[str]
    github: str
    demo: str


class EducationEntry(BaseModel):
    degree: str
    institution: str
    duration: str
    score: str


class SkillCategories(BaseModel):
    languages: list[str]
    frontend: list[str]
    backend_apis: list[str]
    generative_ai: list[str]
    core_cs: list[str]
    tools: list[str]


class ResumeFacts(BaseModel):
    name: str
    tagline: str
    location: str
    email: str
    phone: str
    college: str
    cgpa: str
    about: str
    skills: SkillCategories
    experience: list[ExperienceEntry]
    projects: list[ProjectEntry]
    education: list[EducationEntry]


resume_facts_schema = ResumeFacts.model_json_schema()

# ============================================================
# Project ID generation
# ============================================================

def generate_project_id(project_name: str) -> str:
    """
    Convert a project name into a stable URL/file-system-safe ID.

    Examples:
        Ember Studio
        -> ember-studio

        AI Research Report Generator
        -> ai-research-report-generator

        My Awesome AI Project!
        -> my-awesome-ai-project
    """

    project_id = project_name.lower().strip()

    # Replace anything that isn't a letter or number with a hyphen
    project_id = re.sub(r"[^a-z0-9]+", "-", project_id)

    # Remove leading/trailing hyphens
    project_id = project_id.strip("-")

    return project_id


def assign_project_ids(facts: ResumeFacts) -> ResumeFacts:
    """
    Generate project IDs in Python after Groq extraction.
    """

    for project in facts.projects:
        project.id = generate_project_id(project.name)

    return facts


# ============================================================
# Resume parsing
# ============================================================

def parse_resume(resume_text: str) -> ResumeFacts:

    system_prompt = f"""
You are an expert resume parser.

Extract information from the resume based on its meaning, not only
exact section headings. Different resumes use different headings
(Experience, Professional Experience, Work History, Employment,
Internships) — treat these as equivalent.

Return ONLY valid JSON matching this schema:

{resume_facts_schema}

Important rules:

1. Do not invent information. Only extract what is actually written
   in the resume text below.

2. If a value is not available, return an empty string "" for text
   fields, or an empty list [] for list fields.

3. Include internships inside the "experience" list, not separately.

4. 4. Extract skills mentioned anywhere in the resume — including the
   dedicated skills section, project descriptions, and experience
   descriptions.

   Preserve the categories from the resume and place each skill
   into the closest matching category:

   - languages
   - frontend
   - backend_apis
   - generative_ai
   - core_cs
   - tools

   Do not invent a category.

   If a category has no skills, return an empty list [].

   Example:
   "Languages: C++, Python, JavaScript"
   must be returned as:

   "languages": [
       "C++",
       "Python",
       "JavaScript"
   ]

   "Front-End: React.js, Tailwind CSS"
   must be returned as:

   "frontend": [
       "React.js",
       "Tailwind CSS"
   ]

5. For each project, extract "tech" only if the resume explicitly
   states the tools/stack used for that project. If not mentioned,
   return an empty list for that project's "tech" — do not guess or
   infer technologies from the project description alone.

6. Always return "github": "" and "demo": "" for every project.
   These links are not present in resume text and must not be
   fabricated.

7. Always return "id": "" for every project.
   The project ID will be generated separately by Python after
   extraction.

8. For "education", include every degree/qualification mentioned,
   in the order they appear.
"""

    user_prompt = f"""
Parse the following resume:

{resume_text}
"""

    messages = [
        {
            "role": "system",
            "content": system_prompt
        },
        {
            "role": "user",
            "content": user_prompt
        }
    ]

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format={
            "type": "json_schema",
            "json_schema": {
                "name": "resume_facts",
                "schema": resume_facts_schema
            },
        },
    )

    raw_output = response.choices[0].message.content

    data = json.loads(raw_output)

    facts = ResumeFacts(**data)

    # IMPORTANT:
    # Generate IDs locally instead of trusting the LLM.
    facts = assign_project_ids(facts)

    return facts


# ============================================================
# Save resume facts
# ============================================================

def save_resume_facts(facts: ResumeFacts):
    output_path = (
        Path(__file__).parent
        / "data"
        / "resume_facts.json"
    )

    with open(
        output_path,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            facts.model_dump(),
            file,
            indent=2,
            ensure_ascii=False
        )

    print(
        f"Resume facts saved to: {output_path}"
    )


# ============================================================
# Main
# ============================================================

if __name__ == "__main__":

    resume_path = (
        Path(__file__).parent
        / "data"
        / "source-docs"
        / "resume.docx"
    )

    resume_text = read_docx(resume_path)

    print(
        "Sending to AI for extraction..."
    )

    facts = parse_resume(resume_text)

    save_resume_facts(facts)

    print(
        facts.model_dump_json(indent=2)
    )